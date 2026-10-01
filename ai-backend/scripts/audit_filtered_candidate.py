"""Inspect the 2020 construction workbook without changing its contents."""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def missing_runs(values: list, rows: list) -> list[dict]:
    result = []
    start = None
    for i in range(len(values) + 1):
        missing = i < len(values) and values[i] is None
        if missing and start is None:
            start = i
        if not missing and start is not None:
            result.append({"first_excel_row": rows[start], "last_excel_row": rows[i - 1],
                           "count": i - start})
            start = None
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-plots", action="store_true")
    args = parser.parse_args()
    config = json.loads((ROOT / "configs/datasets.json").read_text())
    dataset = next(d for d in config["datasets"] if d["id"] == "6fd493866k")
    expected = dataset["files"][0]
    source = ROOT / "data/raw/mendeley-6fd493866k-v1" / expected["filename"]
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != expected["sha256"]:
        raise ValueError("Original workbook hash does not match the pinned source")
    workbook = load_workbook(source, data_only=False)
    summary = {"dataset_doi": dataset["doi"], "source_sha256": digest,
               "schema_version": 1, "pollutant": "PM10", "unit": "ug/m3",
               "independent_experiments": 1, "processing": "10-minute moving average",
               "smoothing_alignment": "not documented in workbook",
               "raw_30s_forecast_accepted": False, "sheets": []}
    plot_data = []
    for sheet in workbook:
        entry = {"name": sheet.title, "stored_dimensions": [sheet.max_row, sheet.max_column],
                 "formula_count": sum(c.data_type == "f" for row in sheet for c in row)}
        if sheet.title == "Description of dataset":
            entry["evidence_cells"] = ["B3", "B8", "B14"]
            summary["sheets"].append(entry)
            continue
        if sheet.title not in ("Alpha sensor 10 min moving avg", "Sharp sensor 10 min moving avg"):
            raise ValueError(f"Unexpected measurement sheet: {sheet.title}")
        populated = [(i, row) for i, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3)
                     if any(value is not None for value in row)]
        excel_rows = [i for i, _ in populated]
        rows = [row for _, row in populated]
        clock = np.array([r[0].hour * 3600 + r[0].minute * 60 + r[0].second for r in rows])
        intervals = Counter(np.diff(clock).tolist())
        entry.update({"data_rows": len(rows), "first_excel_row": excel_rows[0],
                      "last_excel_row": excel_rows[-1], "first_clock": str(rows[0][0]),
                      "last_clock": str(rows[-1][0]), "elapsed_seconds": int(clock[-1] - clock[0]),
                      "interval_counts_seconds": {str(k): v for k, v in intervals.items()},
                      "duplicate_clock_count": len(clock) - len(set(clock)),
                      "time_order_increasing": bool(np.all(np.diff(clock) > 0)), "channels": []})
        headers = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        traces = []
        for column, name in enumerate(headers):
            if column == 0 or name is None:
                continue
            original = [r[column] for r in rows]
            invalid = [v for v in original if v is not None and not isinstance(v, (int, float))]
            if invalid:
                raise ValueError(f"Nonnumeric readings in {name}")
            values = np.array([np.nan if v is None else v for v in original], dtype=float)
            valid = values[np.isfinite(values)]
            channel = {"name": name, "observed_count": len(valid),
                       "missing_count": int(np.isnan(values).sum()),
                       "nonfinite_nonmissing_count": int(np.isinf(values).sum()),
                       "negative_count": int((valid < 0).sum()), "zero_count": int((valid == 0).sum()),
                       "unique_observed_values": len(np.unique(valid)),
                       "minimum": float(valid.min()), "maximum": float(valid.max()),
                       "mean": float(valid.mean()), "median": float(np.median(valid)),
                       "p05": float(np.quantile(valid, .05)), "p95": float(np.quantile(valid, .95)),
                       "missing_runs": missing_runs(original, excel_rows)}
            entry["channels"].append(channel)
            traces.append((name, values))
        summary["sheets"].append(entry)
        plot_data.append((sheet.title, (clock - clock[0]) / 60, traces))
    output = ROOT / "reports/data-audit"
    output.mkdir(parents=True, exist_ok=True)
    (output / "mendeley-2020-summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    if not args.skip_plots:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        fig, axes = plt.subplots(2, 1, figsize=(11, 6.5), sharex=True, layout="constrained")
        for axis, (name, elapsed, traces) in zip(axes, plot_data):
            for label, values in traces:
                axis.plot(elapsed, values, label=label, linewidth=1.6)
            axis.set_title(name, loc="left", fontsize=11)
            axis.set_ylabel("PM10 (µg/m³)")
            axis.grid(alpha=.2)
            axis.legend(loc="upper right", fontsize=8)
        axes[-1].set_xlabel("Elapsed minutes from 03:20 (recorded clock; calendar date unspecified)")
        fig.suptitle("Released construction data: ten-minute moving averages", fontsize=14)
        fig.savefig(output / "mendeley-2020-profile.png", dpi=150)
        plt.close(fig)
    print(json.dumps({"summary": str((output / 'mendeley-2020-summary.json').relative_to(ROOT)),
                      "rows_per_measurement_sheet": summary['sheets'][0]['data_rows'],
                      "raw_30s_forecast_accepted": False}))


if __name__ == "__main__":
    main()
